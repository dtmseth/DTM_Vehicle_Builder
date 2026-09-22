"""Project categories preserve Build behavior and make service jobs schedulable."""
from copy import deepcopy
from dataclasses import asdict, replace
from datetime import date

import pytest
from pptx import Presentation

from dtm_buildsheet.app.adapters.wiring import build_local_bundle
from dtm_buildsheet.app.services import agency_service
from dtm_buildsheet.app.services.calendar_service import CalendarService
from dtm_buildsheet.app.services.finalization_service import handle_finalization_check
from dtm_buildsheet.app.services.operations_read_service import OperationsReadService
from dtm_buildsheet.app.services.project_service import handle_save_project, handle_set_project_completion
from dtm_buildsheet.domain.project_codec import project_from_dict
from dtm_buildsheet.domain.project_models import BuildUnit, IndividualUnit
from dtm_buildsheet.domain.project_types import with_project_work
from dtm_buildsheet.domain.operations_models import OperationsActor, VehicleOperations, AcceptanceStatus, VehicleAvailabilityStatus
from dtm_buildsheet.inputs.project_entry import new_project, save_project, load_project
from dtm_buildsheet.inputs.project_drafts import new_draft, save_draft
from dtm_buildsheet.paths import AppPaths

ACTOR = OperationsActor('owner', 'Owner', frozenset({'AppAdmin'}))


@pytest.fixture
def paths(tmp_path):
    for name in ('projects', 'drafts', 'config', 'output', 'agencies'):
        (tmp_path / name).mkdir()
    result = replace(AppPaths(), workspace_dir=tmp_path, workspace_projects_dir=tmp_path/'projects',
                     workspace_drafts_dir=tmp_path/'drafts', workspace_config_dir=tmp_path/'config',
                     workspace_output_dir=tmp_path/'output')
    saved = agency_service.handle_save_agency({"agency_id": "agency", "name": "Agency"}, result)
    assert saved["ok"]
    return result


def create(paths, kind='service', ident='service', **extra):
    body={'project_type': kind, 'customer': {'agency':'Agency', 'agency_id':'agency', 'build_year':'2026'},
          'build_units':[{'unit_id':'u-'+ident, 'vehicle_model':'PIU', 'build_type':'Patrol',
                          'individuals':[{'individual_id':ident, 'unit_number':'101'}]}], **extra}
    result=handle_save_project(body, paths)
    assert result['ok'], result
    return load_project(result['project_id'], paths)


def revision(project):
    return {
        'expected_updated_at': project.updated_at,
        'expected_record_revision': project.record_revision,
    }


def scheduler(paths, project):
    bundle=build_local_bundle()
    ident=project.build_units[0].individuals[0].individual_id
    bundle.operations._records[ident]=VehicleOperations(vehicle_id=ident, project_id=project.project_id,
        acceptance_status=AcceptanceStatus.ACCEPTED, accepted_at='2026-09-01', build_finalized=False,
        parts_status='parts_ready', vehicle_availability_status=VehicleAvailabilityStatus.AT_DTM,
        vehicle_available_date='2026-09-01', parts_received_at='2026-09-01')
    return CalendarService(paths,bundle,cloud=False,clock=lambda:date(2026,9,15))


def preview(service, hours=3.6, **extra):
    view=service.view(ACTOR)
    ident=next(iter(service.bundle.operations._records))
    return service.preview(ACTOR, {'revision':view['revision'],'source_revision':view['source_revision'],
        'edit':{'id':ident,'team_id':'team-david','hours':hours,**extra},'reason':'Service confirmed'})


def test_legacy_defaults_to_build_and_unknown_type_is_rejected(paths):
    project=project_from_dict({'project_id':'old'})
    assert project.project_type=='build'
    assert not handle_save_project({'project_type':'typo'},paths)['ok']


def test_multiple_service_visits_do_not_merge_with_build_or_each_other(paths):
    build=create(paths,'build','build')
    one=create(paths,ident='one');two=create(paths,ident='two')
    assert len({build.project_id,one.project_id,two.project_id})==3
    handle_set_project_completion(build.project_id,{'completed':True},paths)
    done=handle_set_project_completion(one.project_id,{'completed':True},paths)
    assert done['ok'] and not done.get('conflict')
    assert load_project(build.project_id,paths).build_units[0].individuals[0].individual_id=='build'


def test_partial_edit_preserves_type_and_references_without_copying_parts(paths):
    source=create(paths,'build','original')
    service=create(paths)
    link={'project_id':source.project_id,'unit_id':'u-original','individual_id':'original'}
    units=asdict(service)['build_units'];units[0]['individuals'][0]['previous_build']=link
    assert handle_save_project({
        'project_id': service.project_id, 'build_units': units, **revision(service),
    }, paths)['ok']
    service = load_project(service.project_id, paths)
    assert handle_save_project({
        'project_id': service.project_id,
        'project_notes': 'Replace one radio',
        **revision(service),
    }, paths)['ok']
    saved=load_project(service.project_id,paths)
    assert saved.project_type=='service' and saved.build_units[0].individuals[0].previous_build==link
    assert not saved.build_units[0].individuals[0].draft_id
    assert not saved.build_units[0].individuals[0].qb_estimate_id
    assert saved.build_units[0].individuals[0].individual_id=='service'
    units[0]['individuals'][0]['previous_build']['individual_id']='missing'
    assert not handle_save_project({
        'project_id': service.project_id, 'build_units': units, **revision(saved),
    }, paths)['ok']


def test_service_requires_acceptance_and_explicit_hours_but_no_design(paths):
    service=scheduler(paths,create(paths))
    view=service.view(ACTOR)
    assert view['plan']['queue'][0]['kind']=='service'
    assert not view['plan']['queue'][0]['blocked']
    with pytest.raises(ValueError,match='hours'):
        preview(service,None)
    result=preview(service)
    job=result['plan']['jobs'][0]
    assert job['hours']==3.6 and job['end']==job['ready'] and not job['deadline']
    assert not job['finish_segments'] and not job['blocked']
    service.bundle.operations._records['service'].acceptance_status='not_accepted'
    with pytest.raises(ValueError,match='Accept'):
        preview(service)


def test_optional_service_deadline_and_existing_build_default(paths):
    project=create(paths);service=scheduler(paths,project)
    record=service.bundle.operations._records['service']
    assert not preview(service)['plan']['jobs'][0]['deadline']
    record.must_deliver_override_date='2026-10-01'
    assert preview(service)['plan']['jobs'][0]['deadline']=='2026-10-01'
    record.must_deliver_override_date=''
    summary=OperationsReadService(service.bundle.operations).list_vehicle_summaries(ACTOR,projects=[project])['vehicles'][0]
    assert not summary['must_deliver_by_date']
    project.project_type='build';save_project(project,paths)
    assert preview(service)['plan']['jobs'][0]['deadline']=='2026-10-31'


def test_offsite_travel_uses_team_capacity_and_location_readiness(paths):
    project=create(paths,'offsite',service_details={'location':'Agency garage','contact':'Pat 555-0100','travel_hours':3.6,'requires_parts':False})
    service=scheduler(paths,project)
    record=service.bundle.operations._records['service'];record.parts_status='';record.vehicle_availability_status=VehicleAvailabilityStatus.READY_FOR_PICKUP
    first=preview(service)['plan']['jobs'][0]
    assert first['hours']==7.2 and first['end']=='2026-09-15T12:00'
    assert not first['blocked']
    summary=OperationsReadService(service.bundle.operations).list_vehicle_summaries(ACTOR,projects=[project])['vehicles'][0]
    assert summary['ready_to_build'] and summary['parts_and_vehicle_ready']
    assert set(summary['applicable_workstreams'])=={'shop','final_finish'}
    assert record.parts_status=='' and record.vehicle_availability_status=='ready_for_pickup'


def test_service_stripping_and_finishing_are_explicit(paths):
    project=create(paths,service_details={'requires_strip':True,'requires_finishing':True})
    service=scheduler(paths,project)
    job=preview(service)['plan']['jobs'][0]
    assert job['hours']==3.6+service.view(ACTOR)['settings']['teams'][0]['strip_hours']
    assert job['finish_segments'] and job['ready']!=job['end']


def test_service_finalization_does_not_require_warning_lights(paths):
    project=create(paths)
    draft=new_draft();save_draft(draft,paths.workspace_drafts_dir)
    ind=project.build_units[0].individuals[0];ind.draft_id=draft.draft_id
    ind.pdf_path='output/service.pdf';ind.last_exported_at='2999-01-01T00:00:00+00:00';save_project(project,paths)
    check=handle_finalization_check(project.project_id,'u-service','service',paths)
    assert check['ok'] and not check['warnings'] and not check['blocking']
    assert any(c['status']=='not_applicable' for c in check['checks'])
    project.project_type='build';save_project(project,paths)
    assert handle_finalization_check(project.project_id,'u-service','service',paths)['warnings']


def test_service_sheet_has_instructions_and_manifest_without_diagrams(config,paths,monkeypatch):
    from dtm_buildsheet.domain import ProjectInput,PartInput
    from dtm_buildsheet.planning.planner import build_plan
    from dtm_buildsheet import render_ppt
    project=ProjectInput(info={'ProjectID':'service-proof','Agency':'Service Agency','VehicleType':'PIU',
                              'ProjectType':'service','NewVehicle':{'UNIT ID':'101'}},
                         parts=[PartInput(name='Radio',include=True,quantity=1,notes='Replace damaged radio')],
                         notes={'INSTALLATION NOTES':['Diagnose radio fault and replace damaged cable.']})
    plan=build_plan(project,config)
    monkeypatch.setattr(render_ppt,'_load_vehicle_view_config',lambda *a: (_ for _ in ()).throw(AssertionError('Service tried to load diagrams')))
    deck=Presentation(render_ppt.render_plan_to_ppt(plan,replace(config.paths,workspace_output_dir=paths.workspace_output_dir)))
    text='\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert 'SERVICE WORK SHEET' in text and 'Diagnose radio fault' in text and 'radio' in text.lower()
    assert 'Light Heads' not in text and 'SERVICE SCOPE' in text
    assert len(deck.slides)<6


def test_service_visit_files_are_distinct_from_previous_builds(paths):
    from dtm_buildsheet.domain.vehicle_naming import vehicle_folder_name
    from dtm_buildsheet.render_ppt import build_output_filename
    build=create(paths,'build','same');first=create(paths,ident='first');second=create(paths,ident='second')
    names=[vehicle_folder_name(p,p.build_units[0],p.build_units[0].individuals[0]) for p in (build,first,second)]
    assert len(set(names))==3
    info={'Agency':'Agency','VehicleType':'PIU','BuildYear':'2026','NewVehicle':{'UNIT ID':'101'}}
    names=[build_output_filename({**info,'ProjectType':p.project_type,'ProjectID':p.project_id}) for p in (build,first,second)]
    assert len(set(names))==3
